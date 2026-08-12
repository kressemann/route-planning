"""XLSX-Export der aktuell sichtbaren, finalen Stoppreihenfolge.

Reine Formatierung fuer den Ausdruck (Schriftgroesse, Rahmen, Seiteneinrichtung).
Reihenfolge, geloeschte Stopps und manuelle Mengenueberschreibungen kommen
unveraendert aus dem Tourentwurf (``stops``) - dieses Modul aendert daran
nichts, es formatiert nur die Ausgabe.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

COLUMNS = ["Nr", "Address", "Contact", "Kisten", "Schnittware", "Kresse"]

#: Spaltenbreiten - Address/Contact deutlich breiter als die Mengenfelder,
#: damit lange Adressen/Kundennamen (mit Zeilenumbruch) nicht abgeschnitten werden.
_COLUMN_WIDTHS = {
    "Nr": 6,
    "Address": 45,
    "Contact": 28,
    "Kisten": 12,
    "Schnittware": 14,
    "Kresse": 12,
}

#: Spalten, deren Mengenwert bei 0 als leere Zelle statt "0" ausgegeben wird.
_QTY_COLUMNS = {"Kisten", "Schnittware", "Kresse"}

#: Spalten mit linksbuendigem Text (Rest wird zentriert).
_LEFT_ALIGNED_COLUMNS = {"Address", "Contact"}

FONT_SIZE = 16

_THIN_SIDE = Side(style="thin", color="000000")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")


def _blank_if_zero(value: Any) -> Any:
    """Leere Zelle statt '0' fuer Kisten/Schnittware/Kresse - lesbarer Ausdruck."""
    if value == 0:
        return None
    return value


def build_workbook(stops: list[dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tour"

    ws.append(COLUMNS)
    _style_header_row(ws)

    for index, stop in enumerate(stops, start=1):
        ws.append(
            [
                index,
                stop.get("address", ""),
                stop.get("customer_name", ""),
                _blank_if_zero(stop.get("kisten", 0)),
                _blank_if_zero(stop.get("schnittware", 0)),
                _blank_if_zero(stop.get("kresse", 0)),
            ]
        )

    last_row = ws.max_row
    if last_row >= 2:
        _style_data_rows(ws, last_row)

    _set_column_widths(ws)
    _set_print_setup(ws, last_row)

    return wb


def _style_header_row(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = Font(size=FONT_SIZE, bold=True)
        cell.border = _THIN_BORDER
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_data_rows(ws: Worksheet, last_row: int) -> None:
    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(COLUMNS)):
        for cell in row:
            column_name = COLUMNS[cell.column - 1]
            cell.font = Font(size=FONT_SIZE)
            cell.border = _THIN_BORDER
            horizontal = "left" if column_name in _LEFT_ALIGNED_COLUMNS else "center"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


def _set_column_widths(ws: Worksheet) -> None:
    for col_index, name in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=col_index).column_letter
        ws.column_dimensions[letter].width = _COLUMN_WIDTHS[name]


def _set_print_setup(ws: Worksheet, last_row: int) -> None:
    """Querformat, auf eine Seite Breite skaliert (Hoehe frei), Kopfzeile auf
    jeder Druckseite wiederholt, Druckbereich auf die genutzte Tabelle."""
    last_col_letter = ws.cell(row=1, column=len(COLUMNS)).column_letter
    ws.print_area = f"A1:{last_col_letter}{last_row}"
    ws.print_title_rows = "1:1"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3


def workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
