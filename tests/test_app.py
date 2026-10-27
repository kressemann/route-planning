"""End-to-End-Tests gegen die Flask-App mit einem Fake-Zoho-Client.

Es findet zu keinem Zeitpunkt ein echter Netzwerkzugriff auf Zoho statt.
"""

import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from routeplanning.app import create_app
from routeplanning.zoho import EXCLUDED_STATUSES


class FakeZoho:
    """Ersetzt ZohoInventory in Tests - rein lokal, kein Netzwerk.

    Repliziert dieselbe Status-/Datumsfilterung wie der echte Client, damit
    dieser Stand-in dessen Vertrag (nur tatsaechlich zu liefernde Auftraege)
    einhaelt.
    """

    def __init__(self, orders, items, contacts=None):
        self._orders = orders
        self._items = items
        self._contacts = contacts or {}

    def sales_orders_shipped_on(self, day):
        wanted = day.isoformat()
        return [
            o
            for o in self._orders
            if o["shipment_date"] == wanted and o["status"] not in EXCLUDED_STATUSES
        ]

    def get_item(self, item_id):
        return self._items[item_id]

    def get_contact(self, contact_id):
        return self._contacts.get(contact_id, {})


ORDERS = [
    {
        "salesorder_id": "so-1",
        "salesorder_number": "NA-0001",
        "shipment_date": "2026-08-20",
        "status": "confirmed",
        "customer_id": "c1",
        "customer_name": "Kunde A",
        "shipping_address": {"address": "Musterstr. 1", "zip": "12345", "city": "Musterstadt"},
        "line_items": [
            {"item_id": "i-1-9", "name": "Radieschen Mix 1/9", "quantity": 9},
            {"item_id": "i-unknown", "name": "Mystery Artikel", "quantity": 2},
        ],
    },
    {
        "salesorder_id": "so-2",
        "salesorder_number": "NA-0002",
        "shipment_date": "2026-08-20",
        "status": "open",
        "customer_id": "c2",
        "customer_name": "Kunde B",
        "shipping_address": {"address": "Zweite Str. 2", "zip": "54321", "city": "Anderstadt"},
        "line_items": [
            {"item_id": "i-kresse", "name": "Gartenkresse 1/12 Schale/Pad", "quantity": 24},
        ],
    },
    {
        # anderes Datum - darf im Draft vom 2026-08-20 nicht auftauchen
        "salesorder_id": "so-3",
        "salesorder_number": "NA-0003",
        "shipment_date": "2026-08-21",
        "status": "confirmed",
        "customer_id": "c3",
        "customer_name": "Kunde C",
        "shipping_address": {},
        "line_items": [{"item_id": "i-1-9", "name": "Radieschen Mix 1/9", "quantity": 9}],
    },
    {
        # draft-Status - Planungsentwurf, keine echte Lieferung, muss ignoriert werden
        "salesorder_id": "so-4",
        "salesorder_number": "NA-0004",
        "shipment_date": "2026-08-20",
        "status": "draft",
        "customer_id": "c4",
        "customer_name": "Kunde D (Entwurf)",
        "shipping_address": {},
        "line_items": [{"item_id": "i-1-9", "name": "Radieschen Mix 1/9", "quantity": 9}],
    },
]

ITEMS = {
    "i-1-9": {"custom_fields": []},
    "i-unknown": {"custom_fields": []},
    "i-kresse": {"custom_fields": []},
}


@pytest.fixture
def client(tmp_path):
    app = create_app(
        db_path=tmp_path / "test.sqlite3",
        zoho_client_factory=lambda: FakeZoho(ORDERS, ITEMS),
    )
    app.testing = True
    with app.test_client() as c:
        yield c


def test_load_tour_builds_stops_and_warnings(client):
    resp = client.post("/api/tour/load", json={"date": "2026-08-20"})
    assert resp.status_code == 200
    draft = resp.get_json()

    assert len(draft["stops"]) == 2  # c1, c2 - nicht c3 (anderes Datum), nicht c4 (draft-Status)
    by_id = {s["customer_id"]: s for s in draft["stops"]}
    assert by_id["c1"]["kisten"] == 1  # ceil(9/9)
    assert by_id["c2"]["kresse"] == 2  # ceil(24/12)

    assert len(draft["warnings"]) == 1
    assert draft["warnings"][0]["article_name"] == "Mystery Artikel"


def test_load_tour_rejects_bad_date(client):
    resp = client.post("/api/tour/load", json={"date": "not-a-date"})
    assert resp.status_code == 400


def test_get_draft_after_load(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.get("/api/draft/2026-08-20")
    assert resp.status_code == 200
    assert len(resp.get_json()["stops"]) == 2


def test_get_draft_missing_returns_404(client):
    resp = client.get("/api/draft/2026-01-01")
    assert resp.status_code == 404


def test_manual_quantity_override_via_api(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.put(
        "/api/draft/2026-08-20",
        json={
            "stops": [
                {"customer_id": "c2", "kisten": 0, "schnittware": 0, "kresse": 2},
                {"customer_id": "c1", "kisten": 4, "schnittware": 0, "kresse": 0},
            ]
        },
    )
    assert resp.status_code == 200
    draft = resp.get_json()
    ids = [s["customer_id"] for s in draft["stops"]]
    assert ids == ["c2", "c1"]
    by_id = {s["customer_id"]: s for s in draft["stops"]}
    assert by_id["c1"]["kisten"] == 4
    assert by_id["c1"]["kisten_manual"] is True


def test_profile_save_and_apply(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    save = client.post("/api/profiles", json={"name": "Route1", "date": "2026-08-20"})
    assert save.status_code == 200

    # Reihenfolge im Entwurf aendern, dann Profil erneut anwenden -> alte
    # Reihenfolge muss wiederhergestellt werden.
    client.put(
        "/api/draft/2026-08-20",
        json={
            "stops": [
                {"customer_id": "c2", "kisten": 0, "schnittware": 0, "kresse": 2},
                {"customer_id": "c1", "kisten": 1, "schnittware": 0, "kresse": 0},
            ]
        },
    )
    resp = client.post("/api/profiles/Route1/apply", json={"date": "2026-08-20"})
    assert resp.status_code == 200
    ids = [s["customer_id"] for s in resp.get_json()["stops"]]
    assert ids == ["c1", "c2"]

    profiles = client.get("/api/profiles").get_json()["profiles"]
    assert profiles == ["Route1"]


@pytest.mark.parametrize(
    "name",
    [
        "Dienstag MA/HD",  # Regression 27.10.2026: "/" im Namen ergab 404
        "Mo/Di/Mi",
        "A//B",
        "Süd-Tour Größe ÄÖÜ",
        "Tour 50% & mehr?#",
    ],
)
def test_profile_apply_with_special_characters_in_name(client, name):
    from urllib.parse import quote

    client.post("/api/tour/load", json={"date": "2026-08-20"})
    # Profil mit Reihenfolge c1, c2 speichern und eine Fassung mit gleichem
    # Praefix anlegen, damit eine Verwechslung auffallen wuerde.
    assert client.post("/api/profiles", json={"name": name, "date": "2026-08-20"}).status_code == 200
    client.put(
        "/api/draft/2026-08-20",
        json={
            "stops": [
                {"customer_id": "c2", "kisten": 0, "schnittware": 0, "kresse": 2},
                {"customer_id": "c1", "kisten": 1, "schnittware": 0, "kresse": 0},
            ]
        },
    )
    assert client.post("/api/profiles", json={"name": name.split("/")[0] + "X", "date": "2026-08-20"}).status_code == 200

    # Gleiche Kodierung wie das Frontend (encodeURIComponent).
    resp = client.post(
        f"/api/profiles/{quote(name, safe='')}/apply", json={"date": "2026-08-20"}
    )
    assert resp.status_code == 200, resp.get_data(as_text=True)
    stops = resp.get_json()["stops"]
    assert [s["customer_id"] for s in stops] == ["c1", "c2"]
    assert not any(s["not_in_profile"] for s in stops)


def test_profile_apply_unknown_name_with_slash_is_json_error(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.post("/api/profiles/Gibt%2FEs%20nicht/apply", json={"date": "2026-08-20"})
    assert resp.status_code == 400
    assert "nicht gefunden" in resp.get_json()["error"]


def test_delete_stop_removes_it_from_draft(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.delete("/api/draft/2026-08-20/stops/c1")
    assert resp.status_code == 200
    draft = resp.get_json()
    ids = [s["customer_id"] for s in draft["stops"]]
    assert ids == ["c2"]

    # bleibt auch beim erneuten Abrufen des Entwurfs geloescht
    resp = client.get("/api/draft/2026-08-20")
    ids = [s["customer_id"] for s in resp.get_json()["stops"]]
    assert ids == ["c2"]


def test_delete_stop_unknown_returns_404(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.delete("/api/draft/2026-08-20/stops/does-not-exist")
    assert resp.status_code == 404


def test_delete_stop_does_not_appear_in_xlsx(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    client.delete("/api/draft/2026-08-20/stops/c1")

    resp = client.get("/api/draft/2026-08-20/xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    customer_names = [row[2] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "Kunde A" not in customer_names
    assert "Kunde B" in customer_names


def test_xlsx_download_reflects_manual_values(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    client.put(
        "/api/draft/2026-08-20",
        json={
            "stops": [
                {"customer_id": "c1", "kisten": 4, "schnittware": 0, "kresse": 0},
                {"customer_id": "c2", "kisten": 0, "schnittware": 0, "kresse": 2},
            ]
        },
    )
    resp = client.get("/api/draft/2026-08-20/xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == ["Nr", "Address", "Contact", "Kisten", "Schnittware", "Kresse"]
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == 1
    assert row2[2] == "Kunde A"
    assert row2[3] == 4  # manuell ueberschriebener Wert


# -- PDF -----------------------------------------------------------------


def test_pdf_download_content_type_and_filename(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    resp = client.get("/api/draft/2026-08-20/pdf")
    assert resp.status_code == 200
    assert resp.mimetype == "application/pdf"
    assert 'filename="Lieferliste_2026-08-20.pdf"' in resp.headers["Content-Disposition"]


def test_pdf_download_unknown_date_returns_404(client):
    resp = client.get("/api/draft/2026-08-20/pdf")
    assert resp.status_code == 404


def test_delete_stop_does_not_appear_in_pdf(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    client.delete("/api/draft/2026-08-20/stops/c1")

    resp = client.get("/api/draft/2026-08-20/pdf")
    assert resp.status_code == 200
    reader = PdfReader(io.BytesIO(resp.data))
    text = reader.pages[0].extract_text()
    assert "Kunde A" not in text
    assert "Kunde B" in text


def test_pdf_download_reflects_manual_values_and_order(client):
    client.post("/api/tour/load", json={"date": "2026-08-20"})
    client.put(
        "/api/draft/2026-08-20",
        json={
            "stops": [
                {"customer_id": "c2", "kisten": 0, "schnittware": 0, "kresse": 6},
                {"customer_id": "c1", "kisten": 4, "schnittware": 0, "kresse": 0},
            ]
        },
    )
    resp = client.get("/api/draft/2026-08-20/pdf")
    assert resp.status_code == 200

    reader = PdfReader(io.BytesIO(resp.data))
    text = reader.pages[0].extract_text()
    assert "4" in text  # manuell ueberschriebener Kisten-Wert fuer Kunde A
    assert "6" in text  # manuell ueberschriebener Kresse-Wert fuer Kunde B
    # Reihenfolge aus dem Tourentwurf (c2 vor c1) muss uebernommen werden.
    assert text.index("Kunde B") < text.index("Kunde A")
